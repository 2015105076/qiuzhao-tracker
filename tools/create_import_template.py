"""Create the Excel import template shipped in assets/import-template.xlsx."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "assets" / "import-template.xlsx"

ACCENT = "C24E35"
ACCENT_DARK = "A53B26"
INK = "201E19"
INK_SOFT = "49463E"
MUTED = "746D5F"
PAPER = "F4F0E7"
PAPER_LIGHT = "FBF8F1"
PAPER_DEEP = "EBE5D8"
LINE = "DDD5C5"
GREEN = "4E7157"
GOLD = "9C782F"

FONT_NAME = "Microsoft YaHei"
THIN = Side(style="thin", color=LINE)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(cell):
    cell.fill = PatternFill("solid", fgColor=ACCENT)
    cell.font = Font(name=FONT_NAME, size=11, bold=True, color="FFF8EC")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = BORDER


def style_section(cell):
    cell.fill = PatternFill("solid", fgColor=INK)
    cell.font = Font(name=FONT_NAME, size=10, bold=True, color="FFF8EC")
    cell.alignment = Alignment(horizontal="left", vertical="center")


def style_body(cell, color=INK_SOFT, bold=False):
    cell.font = Font(name=FONT_NAME, size=10, color=color, bold=bold)
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    cell.border = BORDER


workbook = Workbook()
workbook.properties.title = "投递簿 · Excel 批量导入模板"
workbook.properties.subject = "秋招投递记录批量导入"
workbook.properties.creator = "投递簿"
workbook.properties.description = "在“投递记录”工作表中填写，然后上传到投递簿。"
workbook.properties.keywords = "秋招,投递,求职,模板"

# ---------------------------------------------------------------------------
# Data sheet
# ---------------------------------------------------------------------------
sheet = workbook.active
sheet.title = "投递记录"
sheet.sheet_properties.tabColor = ACCENT

headers = ["企业", "岗位", "投递时间", "状态", "投递渠道", "工作城市", "备注"]
widths = [22, 28, 16, 14, 16, 14, 42]
comments = [
    "必填。公司或单位名称，最多 60 字。\n示例：字节跳动",
    "必填。投递的岗位名称，最多 80 字。\n示例：产品经理（校招）",
    "必填。推荐 YYYY-MM-DD，也可直接使用 Excel 日期格式。\n示例：2026-10-24",
    "选填。不填时默认按“已投递”导入。\n可选：已投递、筛选中、笔试、面试中、已获 Offer、未通过",
    "选填。可选：官网投递、内推、招聘平台、校园宣讲、其他。\n无法识别时会归入“其他”。",
    "选填。工作城市，最多 40 字。\n示例：上海",
    "选填。面试安排、内推人、需要准备的材料等，最多 200 字。",
]
for column, (title, width, hint) in enumerate(zip(headers, widths, comments), start=1):
    cell = sheet.cell(row=1, column=column, value=title)
    style_header(cell)
    cell.comment = Comment(hint, "投递簿", width=260, height=95)
    sheet.column_dimensions[cell.column_letter].width = width

sheet.row_dimensions[1].height = 31
sheet.freeze_panes = "A2"
sheet.auto_filter.ref = "A1:G1"
sheet.sheet_view.zoomScale = 100

# Give the first 200 rows a clean, readable writing area.
for row in range(2, 202):
    fill = PatternFill("solid", fgColor=PAPER_LIGHT if row % 2 == 0 else "FFFFFF")
    for column in range(1, 8):
        cell = sheet.cell(row=row, column=column)
        cell.fill = fill
        cell.border = BORDER
        cell.font = Font(name=FONT_NAME, size=10, color=INK_SOFT)
        cell.alignment = Alignment(vertical="center")
    sheet.cell(row=row, column=3).number_format = "yyyy-mm-dd"
    sheet.row_dimensions[row].height = 23

status_validation = DataValidation(
    type="list",
    formula1='"已投递,筛选中,笔试,面试中,已获 Offer,未通过"',
    allow_blank=True,
    showErrorMessage=True,
    errorTitle="状态不在可选范围内",
    error="请从下拉列表中选择状态，或留空按“已投递”导入。",
    promptTitle="当前状态",
    prompt="选择这条投递的当前状态。",
)
channel_validation = DataValidation(
    type="list",
    formula1='"官网投递,内推,招聘平台,校园宣讲,其他"',
    allow_blank=True,
    showErrorMessage=True,
    errorTitle="渠道不在可选范围内",
    error="请从下拉列表中选择投递渠道。",
    promptTitle="投递渠道",
    prompt="选择这条投递的投递渠道。",
)
sheet.add_data_validation(status_validation)
status_validation.add("D2:D1001")
sheet.add_data_validation(channel_validation)
channel_validation.add("E2:E1001")

# ---------------------------------------------------------------------------
# Guidance sheet
# ---------------------------------------------------------------------------
guide = workbook.create_sheet("填写说明")
guide.sheet_properties.tabColor = GREEN
for letter, width in {"A": 16, "B": 14, "C": 52, "D": 46}.items():
    guide.column_dimensions[letter].width = width

guide.merge_cells("A1:D1")
title = guide["A1"]
title.value = "投递簿 · Excel 批量导入说明"
title.fill = PatternFill("solid", fgColor=PAPER)
title.font = Font(name=FONT_NAME, size=20, bold=True, color=INK)
title.alignment = Alignment(vertical="center")
guide.row_dimensions[1].height = 46

guide.merge_cells("A2:D2")
subtitle = guide["A2"]
subtitle.value = "在“投递记录”工作表中填写数据，然后回到投递簿上传文件。导入前会先预览校验，不会直接写入。"
subtitle.fill = PatternFill("solid", fgColor=PAPER)
subtitle.font = Font(name=FONT_NAME, size=10, color=MUTED)
subtitle.alignment = Alignment(vertical="center")
guide.row_dimensions[2].height = 28

guide["A4"] = "填写步骤"
guide.merge_cells("A4:D4")
style_section(guide["A4"])
guide.row_dimensions[4].height = 26

steps = [
    ("1", "打开“投递记录”工作表，从第 2 行开始逐行填写；第一行字段名请勿修改或删除。"),
    ("2", "企业、岗位、投递时间为必填项；状态不填时默认按“已投递”导入。"),
    ("3", "每一行代表一条投递记录。空行会自动跳过，不需要手动删除。"),
    ("4", "回到投递簿，点击“导入 Excel”，选择保存好的文件并检查预览结果。"),
]
for index, (number, text) in enumerate(steps, start=5):
    guide.cell(row=index, column=1, value=number)
    guide.merge_cells(start_row=index, start_column=2, end_row=index, end_column=4)
    guide.cell(row=index, column=2, value=text)
    for column in range(1, 5):
        cell = guide.cell(row=index, column=column)
        cell.fill = PatternFill("solid", fgColor="FFFFFF")
        cell.border = BORDER
        cell.font = Font(name=FONT_NAME, size=10, color=INK_SOFT)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    guide.cell(row=index, column=1).alignment = Alignment(horizontal="center", vertical="center")
    guide.row_dimensions[index].height = 27

field_header_row = 10
field_headers = ["字段", "是否必填", "填写说明", "可识别表头 / 示例"]
for column, value in enumerate(field_headers, start=1):
    cell = guide.cell(row=field_header_row, column=column, value=value)
    style_header(cell)
guide.row_dimensions[field_header_row].height = 27

fields = [
    ("企业", "必填", "公司或单位名称，最多 60 字。", "企业、公司、企业名称、公司名称、company"),
    ("岗位", "必填", "投递的岗位名称，最多 80 字。", "岗位、职位、投递岗位、应聘岗位、position"),
    ("投递时间", "必填", "推荐 YYYY-MM-DD；也支持 Excel 日期、2026/10/24、2026年10月24日。", "投递时间、投递日期、申请日期、日期、date"),
    ("状态", "选填", "留空时按“已投递”导入。", "已投递、筛选中、笔试、面试中、已获 Offer、未通过"),
    ("投递渠道", "选填", "下拉选择，无法识别时归入“其他”。", "官网投递、内推、招聘平台、校园宣讲、其他"),
    ("工作城市", "选填", "城市名称，最多 40 字。", "城市、工作城市、地点、location"),
    ("备注", "选填", "内推人、面试安排等，最多 200 字。", "备注、说明、notes"),
]
for row_index, row_data in enumerate(fields, start=field_header_row + 1):
    for column, value in enumerate(row_data, start=1):
        cell = guide.cell(row=row_index, column=column, value=value)
        style_body(cell, color=ACCENT if column == 2 and value == "必填" else INK_SOFT, bold=column == 1)
    guide.row_dimensions[row_index].height = 29

status_row = 19
guide.cell(row=status_row, column=1, value="状态说明")
guide.merge_cells(start_row=status_row, start_column=1, end_row=status_row, end_column=4)
style_section(guide.cell(row=status_row, column=1))
guide.row_dimensions[status_row].height = 26

statuses = [
    ("已投递", "简历已提交，等待企业处理。"),
    ("筛选中", "企业正在筛选简历或安排后续流程。"),
    ("笔试", "已进入笔试、测评或在线测试环节。"),
    ("面试中", "正在面试流程中，包括一面、二面、HR 面等。"),
    ("已获 Offer", "已经收到录用通知或意向书。"),
    ("未通过", "简历未通过、流程终止或已明确拒绝。"),
]
for row_index, (status, description) in enumerate(statuses, start=status_row + 1):
    guide.cell(row=row_index, column=1, value=status)
    guide.merge_cells(start_row=row_index, start_column=2, end_row=row_index, end_column=4)
    guide.cell(row=row_index, column=2, value=description)
    for column in range(1, 5):
        cell = guide.cell(row=row_index, column=column)
        cell.fill = PatternFill("solid", fgColor="FFFFFF")
        cell.border = BORDER
        cell.font = Font(name=FONT_NAME, size=10, color=INK_SOFT, bold=column == 1)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    guide.row_dimensions[row_index].height = 23

example_row = 27
guide.cell(row=example_row, column=1, value="填写示例")
guide.merge_cells(start_row=example_row, start_column=1, end_row=example_row, end_column=4)
style_section(guide.cell(row=example_row, column=1))
guide.row_dimensions[example_row].height = 26

example_text = (
    "字节跳动 · 产品经理（校招） · 2026-10-24 · 面试中 · 内推 · 上海 · "
    "一面通过，等待二面安排。"
)
guide.merge_cells(start_row=example_row + 1, start_column=1, end_row=example_row + 1, end_column=4)
example = guide.cell(row=example_row + 1, column=1, value=example_text)
example.fill = PatternFill("solid", fgColor=PAPER_LIGHT)
example.font = Font(name=FONT_NAME, size=10, color=INK_SOFT)
example.border = BORDER
example.alignment = Alignment(vertical="center", wrap_text=True)
guide.row_dimensions[example_row + 1].height = 38

guide.merge_cells("A30:D30")
note = guide["A30"]
note.value = "提示：填错状态或日期时，导入预览会标出具体行号，不会写入该行。修正后重新上传即可。"
note.fill = PatternFill("solid", fgColor="F2DED4")
note.font = Font(name=FONT_NAME, size=10, color=ACCENT_DARK)
note.alignment = Alignment(vertical="center", wrap_text=True)
guide.row_dimensions[30].height = 29

guide.sheet_view.showGridLines = False
guide.sheet_view.zoomScale = 100
guide.freeze_panes = "A3"

sheet.sheet_view.selection[0].activeCell = "A2"
sheet.sheet_view.selection[0].sqref = "A2"
guide.sheet_view.selection[0].activeCell = "A1"
guide.sheet_view.selection[0].sqref = "A1"

workbook.active = 0
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
workbook.save(OUTPUT)
print(f"saved {OUTPUT} ({OUTPUT.stat().st_size} bytes)")
